from __future__ import annotations

import importlib
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from io import BytesIO

import cv2
import numpy as np
from openpyxl import load_workbook
from fastapi.testclient import TestClient


class WebApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        os.environ["PARKING_WEB_DATA"] = cls.temp.name
        sys.modules.pop("parking_web.api", None)
        cls.module = importlib.import_module("parking_web.api")
        cls.client = TestClient(cls.module.app)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()
        os.environ.pop("PARKING_WEB_DATA", None)

    def _video_bytes(self) -> bytes:
        path = Path(self.temp.name) / "upload.mp4"
        writer = cv2.VideoWriter(str(path), cv2.VideoWriter_fourcc(*"mp4v"), 10, (32, 24))
        for index in range(5):
            writer.write(np.full((24, 32, 3), index * 20, dtype=np.uint8))
        writer.release()
        return path.read_bytes()

    def test_00_settings_persist_and_validate(self):
        original = self.client.get("/api/settings").json()
        self.assertEqual(original["evaluation"]["min_temporal_iou"], .30)
        changed = {**original,
                   "run_defaults": {**original["run_defaults"], "plate_top_k": 7,
                                    "grace_seconds": 2.5},
                   "evaluation": {**original["evaluation"], "min_temporal_iou": .6}}
        try:
            response = self.client.put("/api/settings", json=changed)
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(self.client.get("/api/settings").json()["run_defaults"]["plate_top_k"], 7)
            self.assertEqual(self.module.storage.get_settings()["evaluation"]["min_temporal_iou"], .6)
            invalid = self.client.put("/api/settings", json={**changed,
                "run_defaults": {**changed["run_defaults"], "condition_max_frames": 6}})
            self.assertEqual(invalid.status_code, 422)
            self.assertEqual(self.client.get("/api/settings").json()["run_defaults"]["condition_max_frames"], 3)
        finally:
            self.client.put("/api/settings", json=original)

    def test_01_upload_preview_dashboard_and_config_validation(self):
        response = self.client.post("/api/videos", content=self._video_bytes(),
                                    headers={"Content-Type": "application/octet-stream",
                                             "X-Filename": "camera%20entry.mp4"})
        self.assertEqual(response.status_code, 201, response.text)
        video = response.json()
        self.assertEqual(video["filename"], "camera entry.mp4")
        self.assertEqual(self.client.get(f"/api/videos/{video['id']}/frame").status_code, 200)
        dashboard = self.client.get("/api/dashboard").json()
        self.assertEqual(dashboard["video_count"], 1)
        payload = {"video_id": video["id"], "camera": {
            "camera_id": "cam", "lane_id": "lane", "direction": "ENTRY",
            "roi": [[.2, .2], [.8, .2], [.8, .8], [.2, .8]],
            "crossing_line": [[.1, .05], [.9, .05]],
        }}
        invalid = self.client.post("/api/configs/validate", json=payload)
        self.assertEqual(invalid.status_code, 422)
        payload["camera"]["crossing_line"] = [[.1, .5], [.9, .5]]
        self.assertEqual(self.client.post("/api/configs/validate", json=payload).status_code, 200)
        advanced = {**payload, "plate_confidence": .42, "plate_top_k": 7,
                    "plate_min_gap_ms": 450, "plate_class_name": "Vehicle registration plate",
                    "motion_gate_enabled": True, "motion_gate_probe_seconds": .5}
        config = self.module._config(self.module.RunPayload.model_validate(advanced),
                                     Path(self.temp.name) / "validation")
        self.assertEqual((config.plate_confidence, config.plate_top_k,
                          config.plate_min_gap_ms), (.42, 7, 450))
        self.assertTrue(config.motion_gate_enabled)
        self.assertEqual(config.motion_gate_probe_seconds, .5)
        self.assertEqual(self.client.post("/api/configs/validate", json={**advanced,
            "motion_gate_probe_seconds": 0}).status_code, 422)

    def test_02_result_annotation_and_review_flow(self):
        video = self.module.storage.list_videos()[0]
        config = {"video_id": video["id"], "camera": {"camera_id": "cam", "lane_id": "lane",
                  "direction": "ENTRY", "roi": [], "crossing_line": []}}
        web_run = self.module.storage.create_run(video["id"], config)
        run_dir = Path(self.temp.name) / "synthetic-run"
        (run_dir / "clips").mkdir(parents=True)
        event = {"event_id": "EVT-000001", "track_id": 1, "vehicle_type": "car",
                 "start_ms": 100, "end_ms": 500, "crossed": True, "clip_id": "CLIP-000001",
                 "camera_id": "cam", "lane_id": "lane", "direction": "ENTRY"}
        clip = {"clip_id": "CLIP-000001", "start_ms": 0, "end_ms": 800,
                "actual_start_ms": 0, "clip_uri": "clips/CLIP-000001.mp4"}
        (run_dir / "events.jsonl").write_text(json.dumps(event) + "\n", encoding="utf-8")
        (run_dir / "clips.jsonl").write_text(json.dumps(clip) + "\n", encoding="utf-8")
        source_capture = cv2.VideoCapture(video["stored_path"])
        clip_path = run_dir / clip["clip_uri"]
        clip_writer = cv2.VideoWriter(str(clip_path), cv2.VideoWriter_fourcc(*"mp4v"), 10, (32, 24))
        while True:
            ok, frame = source_capture.read()
            if not ok:
                break
            clip_writer.write(frame)
        clip_writer.release()
        source_capture.release()
        (run_dir / "detections.jsonl").write_text("", encoding="utf-8")
        plate = {"event_id": "EVT-000001", "observations": [{"timestamp_ms": 200,
                 "plate_box": [10, 20, 30, 40], "plate_confidence": 0.8}],
                 "consensus": {"text": "ABC123", "confidence": 0.9, "support_count": 2}}
        (run_dir / "plate_observations.jsonl").write_text(json.dumps(plate) + "\n", encoding="utf-8")
        condition = {"event_id": "EVT-000001", "plate_readable": True,
                     "condition_status": "good", "vlm_readable": True,
                     "cv_metrics": [{"suggested_conditions": []},
                                    {"suggested_conditions": ["capture_blur"]}]}
        (run_dir / "condition_suggestions.jsonl").write_text(
            json.dumps(condition) + "\n", encoding="utf-8")
        (run_dir / "run.json").write_text(json.dumps({"artifacts": {"final_clip": None}}), encoding="utf-8")
        self.module.storage.update_run(web_run["id"], status="COMPLETED", progress=100,
                                       run_dir=str(run_dir), pipeline_run_id="step1-test",
                                       event_count=1, clip_count=1)
        result = self.client.get(f"/api/runs/{web_run['id']}/result")
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()["events"][0]["clip"]["clip_id"], "CLIP-000001")
        self.assertEqual(result.json()["events"][0]["plate_observations"][0]["plate_box"], [10, 20, 30, 40])
        suggestion = result.json()["events"][0]["suggestion"]
        self.assertEqual(suggestion["plate_consensus_confidence"], 0.9)
        self.assertEqual(suggestion["plate_support_count"], 2)
        self.assertEqual(suggestion["cv_clean_fraction"], 0.5)
        self.assertTrue(suggestion["vlm_readable"])
        best_url = f"/api/runs/{web_run['id']}/events/EVT-000001/best-frame"
        self.assertEqual(self.client.get(best_url, params={"start_ms": 100, "end_ms": 500}).json(),
                         {"timestamp_ms": 200})
        self.assertEqual(self.client.get(best_url, params={"start_ms": 300, "end_ms": 500}).json(),
                         {"timestamp_ms": 400})
        manual_url = f"/api/runs/{web_run['id']}/events/GT-000001/best-frame"
        self.assertEqual(self.client.get(manual_url, params={"start_ms": 100, "end_ms": 500}).json(),
                         {"timestamp_ms": 300})
        self.assertEqual(self.client.get(best_url, params={"start_ms": 500, "end_ms": 100}).status_code, 422)
        self.assertEqual(self.client.get(f"/api/runs/{web_run['id']}/events/UNKNOWN/best-frame",
                                         params={"start_ms": 100, "end_ms": 500}).status_code, 404)
        (run_dir / "condition_suggestions.jsonl").unlink()
        media = self.client.get(f"/api/runs/{web_run['id']}/media/{clip['clip_uri']}",
                                headers={"Range": "bytes=0-31"})
        self.assertEqual(media.status_code, 206, media.text)
        self.assertEqual(media.headers["content-type"], "video/mp4")
        self.assertEqual(int(media.headers["content-range"].split("/")[1]), clip_path.stat().st_size)
        preview = self.client.get(f"/api/runs/{web_run['id']}/media/{clip['clip_uri']}?preview=true",
                                  headers={"Range": "bytes=0-31"})
        self.assertEqual(preview.status_code, 206, preview.text)
        self.assertEqual(preview.headers["content-type"], "video/mp4")
        cached_files = list((run_dir / ".browser_media").glob("*.mp4"))
        self.assertEqual(len(cached_files), 1)
        browser_clip = cv2.VideoCapture(str(cached_files[0]))
        codec = int(browser_clip.get(cv2.CAP_PROP_FOURCC))
        self.assertEqual("".join(chr((codec >> (8 * i)) & 255) for i in range(4)).lower(), "h264")
        browser_clip.release()
        annotation = {"event_id": "EVT-000001", "vehicle_type": "car", "crossed": True,
                      "plate_text": "59A12345", "plate_readable": True,
                      "condition_status": "good", "conditions": [], "start_ms": 100, "end_ms": 500}
        export = self.client.post(f"/api/runs/{web_run['id']}/export.xlsx", json=[annotation])
        self.assertEqual(export.status_code, 200, export.text)
        book = load_workbook(BytesIO(export.content))
        sheet = book.active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet["A2"].value, "EVT-000001")
        self.assertEqual(sheet["F2"].value, "59A12345")
        self.assertEqual(len(sheet._images), 1)
        self.assertNotIn("Detector TB", [cell.value for cell in sheet[1]])
        self.assertNotIn("Độ dễ (0–1)", [cell.value for cell in sheet[1]])
        saved = self.client.put(f"/api/runs/{web_run['id']}/annotations", json=[annotation])
        self.assertEqual(saved.status_code, 200, saved.text)
        self.assertTrue((run_dir / "ground_truth.jsonl").is_file())
        reviewed = self.client.post(f"/api/runs/{web_run['id']}/review")
        self.assertEqual(reviewed.status_code, 200, reviewed.text)
        self.assertEqual(reviewed.json()["review_status"], "REVIEWED")
        unchanged = self.client.put(f"/api/runs/{web_run['id']}/annotations", json=[annotation])
        self.assertEqual(unchanged.status_code, 200, unchanged.text)
        self.assertEqual(unchanged.json()["review_status"], "REVIEWED")
        self.assertEqual(self.client.get(f"/api/runs/{web_run['id']}").json()["review_status"], "REVIEWED")
        self.assertEqual(self.client.get(f"/api/runs/{web_run['id']}/result").json()["run"]["review_status"], "REVIEWED")
        invalid = self.client.put(f"/api/runs/{web_run['id']}/annotations",
                                  json=[{**annotation, "end_ms": 0}])
        self.assertEqual(invalid.status_code, 400)
        self.assertEqual(self.client.get(f"/api/runs/{web_run['id']}").json()["review_status"], "REVIEWED")

        changed = self.client.put(f"/api/runs/{web_run['id']}/annotations",
                                  json=[{**annotation, "plate_text": "59A12346"}])
        self.assertEqual(changed.status_code, 200, changed.text)
        self.assertEqual(changed.json()["review_status"], "DRAFT")
        self.assertEqual(self.client.post(f"/api/runs/{web_run['id']}/review").status_code, 200)

        evaluation = self.client.get(f"/api/runs/{web_run['id']}/evaluation")
        self.assertEqual(evaluation.status_code, 200, evaluation.text)
        self.assertEqual(evaluation.json()["runtime"]["scope"], "unavailable")
        (run_dir / "runtime.json").write_text(json.dumps({"schema_version": "0.2.0",
            "scope": "full_pipeline", "elapsed_seconds": 12.5,
            "core_elapsed_seconds": 3.0, "ocr_elapsed_seconds": 4.0,
            "condition_elapsed_seconds": 5.5}), encoding="utf-8")
        timed = self.client.get(f"/api/runs/{web_run['id']}/evaluation").json()
        self.assertEqual(timed["runtime"]["elapsed_seconds"], 12.5)
        self.assertEqual(timed["runtime"]["scope"], "full_pipeline")
        self.assertEqual(timed["runtime"]["core_elapsed_seconds"], 3.0)
        self.assertEqual(timed["runtime"]["ocr_elapsed_seconds"], 4.0)
        self.assertEqual(timed["runtime"]["condition_elapsed_seconds"], 5.5)
        self.assertAlmostEqual(timed["runtime"]["total_to_raw_ratio"],
                               12.5 / timed["runtime"]["raw_duration_seconds"])
        self.assertIsNone(timed["runtime"]["total_processing_seconds"])
        self.assertIsNone(timed["runtime"]["processing_to_raw_ratio"])
        (run_dir / "runtime.json").write_text(json.dumps({"schema_version": "0.3.0",
            "scope": "full_pipeline", "elapsed_seconds": 12.5,
            "core_elapsed_seconds": 3.0, "ocr_elapsed_seconds": 4.0,
            "condition_elapsed_seconds": 5.5,
            "core_profile_seconds": {"detector_load_seconds": 0.5},
            "ocr_profile_seconds": {"model_load_seconds": 1.0},
            "condition_profile_seconds": {"model_load_seconds": 1.5}}), encoding="utf-8")
        profiled = self.client.get(f"/api/runs/{web_run['id']}/evaluation").json()["runtime"]
        self.assertEqual(profiled["core_processing_seconds"], 2.5)
        self.assertEqual(profiled["ocr_processing_seconds"], 3.0)
        self.assertEqual(profiled["condition_processing_seconds"], 4.0)
        self.assertEqual(profiled["total_processing_seconds"], 9.5)
        self.assertAlmostEqual(profiled["processing_to_raw_ratio"],
                               9.5 / profiled["raw_duration_seconds"])
        (run_dir / "runtime.json").unlink()
        (run_dir / "run.json").write_text(json.dumps({"artifacts": {"final_clip": None},
            "elapsed_seconds": 8.5}), encoding="utf-8")
        legacy = self.client.get(f"/api/runs/{web_run['id']}/evaluation").json()
        self.assertEqual(legacy["runtime"]["scope"], "step1_core_only")
        self.assertEqual(legacy["runtime"]["elapsed_seconds"], 8.5)
        self.assertEqual(legacy["runtime"]["core_elapsed_seconds"], 8.5)
        self.assertIsNone(legacy["runtime"]["ocr_elapsed_seconds"])
        self.assertIsNone(legacy["runtime"]["total_to_raw_ratio"])
        self.assertEqual(evaluation.json()["metrics"]["events"]["recall"], 1)
        self.assertEqual(evaluation.json()["metrics"]["ocr"]["fn"], 1)
        self.assertEqual(evaluation.json()["metrics"]["clip"]["retention_recall"], 1)
        self.assertFalse(evaluation.json()["metrics"]["condition_status"]["available"])
        original_settings = self.client.get("/api/settings").json()
        try:
            changed_settings = {**original_settings, "evaluation": {
                **original_settings["evaluation"], "min_temporal_iou": .6,
                "coverage_threshold": .8}}
            self.assertEqual(self.client.put("/api/settings", json=changed_settings).status_code, 200)
            rescored = self.client.get(f"/api/runs/{web_run['id']}/evaluation").json()["metrics"]
            self.assertEqual(rescored["events"]["min_temporal_iou"], .6)
            self.assertEqual(rescored["clip"]["coverage_threshold"], .8)
        finally:
            self.client.put("/api/settings", json=original_settings)

        invalid_condition = self.client.put(f"/api/runs/{web_run['id']}/annotations",
                                            json=[{**annotation, "condition_status": "not_applicable"}])
        self.assertEqual(invalid_condition.status_code, 400)
        self.assertEqual(self.client.get(f"/api/runs/{web_run['id']}").json()["review_status"], "REVIEWED")

        self.assertEqual(self.client.post(f"/api/runs/{web_run['id']}/reopen").status_code, 200)
        self.assertEqual(self.client.get(f"/api/runs/{web_run['id']}/evaluation").status_code, 409)
        manual = {**annotation, "event_id": "GT-000001", "start_ms": 0,
                  "end_ms": 80, "plate_text": None, "plate_readable": False}
        saved = self.client.put(f"/api/runs/{web_run['id']}/annotations", json=[annotation, manual])
        self.assertEqual(saved.status_code, 200, saved.text)
        result = self.client.get(f"/api/runs/{web_run['id']}/result").json()
        self.assertEqual(len(result["gt_events"]), 2)
        self.assertIsNone(next(x for x in result["gt_events"] if x["event_id"] == "GT-000001")["source_event_id"])
        self.assertEqual(self.client.post(f"/api/runs/{web_run['id']}/review").status_code, 200)
        scores = self.client.get(f"/api/runs/{web_run['id']}/evaluation").json()["metrics"]["events"]
        self.assertEqual((scores["tp"], scores["fp"], scores["fn"]), (1, 0, 1))
        self.assertEqual(scores["recall"], .5)

        self.assertEqual(self.client.post(f"/api/runs/{web_run['id']}/reopen").status_code, 200)
        rejected = {**annotation, "is_valid_event": False}
        self.assertEqual(self.client.put(f"/api/runs/{web_run['id']}/annotations", json=[rejected]).status_code, 200)
        self.assertEqual(self.client.post(f"/api/runs/{web_run['id']}/review").status_code, 200)
        self.assertEqual((run_dir / "ground_truth.jsonl").read_text(encoding="utf-8"), "")
        self.assertEqual(self.client.get(f"/api/runs/{web_run['id']}/result").json()["gt_events"], [])
        scores = self.client.get(f"/api/runs/{web_run['id']}/evaluation").json()["metrics"]["events"]
        self.assertEqual((scores["tp"], scores["fp"], scores["fn"]), (0, 1, 0))

    def test_03_delete_run_removes_artifacts_and_gt_but_keeps_raw_video(self):
        video = self.module.storage.list_videos()[0]
        web_run = self.module.storage.create_run(video["id"], {"video_id": video["id"]})
        run_id = web_run["id"]
        self.assertEqual(self.client.delete(f"/api/runs/{run_id}").status_code, 409)
        self.assertIsNotNone(self.module.storage.get_run(run_id))

        run_dir = self.module.storage.runs / "delete-test"
        (run_dir / "clips").mkdir(parents=True)
        (run_dir / "clips" / "clip.mp4").write_bytes(b"clip")
        (run_dir / "events.jsonl").write_text("{}\n", encoding="utf-8")
        (run_dir / ".browser_media").mkdir()
        (run_dir / ".browser_media" / "cached.mp4").write_bytes(b"cached")
        self.module.storage.update_run(run_id, status="COMPLETED", run_dir=str(run_dir),
                                       review_status="REVIEWED", event_count=1)
        self.module.storage.save_annotation(run_id, "EVT-000001", {"event_id": "EVT-000001"})
        before = self.module.storage.dashboard()
        deleted = self.client.delete(f"/api/runs/{run_id}")
        self.assertEqual(deleted.status_code, 200, deleted.text)
        self.assertFalse(run_dir.exists())
        self.assertIsNone(self.module.storage.get_run(run_id))
        self.assertEqual(self.module.storage.annotations(run_id), {})
        self.assertEqual(self.client.get(f"/api/runs/{run_id}").status_code, 404)
        self.assertEqual(self.client.delete(f"/api/runs/{run_id}").status_code, 404)
        self.assertEqual(self.module.storage.dashboard()["run_count"], before["run_count"] - 1)
        self.assertTrue(Path(video["stored_path"]).is_file())

        external = Path(self.temp.name) / "not-a-run"
        external.mkdir()
        (external / "keep.txt").write_text("keep", encoding="utf-8")
        unsafe = self.module.storage.create_run(video["id"], {"video_id": video["id"]})
        self.module.storage.update_run(unsafe["id"], status="COMPLETED", run_dir=str(external))
        self.assertEqual(self.client.delete(f"/api/runs/{unsafe['id']}").status_code, 409)
        self.assertTrue((external / "keep.txt").is_file())


if __name__ == "__main__":
    unittest.main()
