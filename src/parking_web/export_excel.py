"""Export the current review table with a traceable annotated raw frame."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import cv2
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill

from parking_step1.evaluation import load_jsonl


HEADERS = ("Event", "Bắt đầu ms", "Kết thúc ms", "Loại", "Crossed",
           "Biển số", "Đọc biển", "Trạng thái", "Điều kiện", "Clip",
           "Video", "Frame đại diện")


def _best_frame(event: dict | None, start_ms: int, end_ms: int,
                plate: dict | None, detections: list[dict]) -> tuple[int, list | None, list | None]:
    if event:
        observations = [item for item in (plate or {}).get("observations", [])
                        if start_ms <= item.get("timestamp_ms", -1) <= end_ms]
        if observations:
            best = max(observations, key=lambda item: (
                .45 * float(item.get("plate_confidence") or 0) +
                .40 * float((item.get("quality") or {}).get("score") or 0) +
                .15 * float(item.get("ocr_confidence") or 0),
                -item["timestamp_ms"]))
            return best["timestamp_ms"], best.get("vehicle_box"), best.get("plate_box")
        candidates = [(record["timestamp_ms"], detection) for record in detections
                      if start_ms <= record["timestamp_ms"] <= end_ms
                      for detection in record.get("detections", [])
                      if detection.get("track_id") == event.get("track_id")]
        if candidates:
            timestamp, detection = max(candidates, key=lambda item: (
                float(item[1].get("confidence") or 0), -item[0]))
            return timestamp, detection.get("box"), None
    return start_ms + max(0, end_ms - start_ms) // 2, None, None


def _annotated_image(capture: cv2.VideoCapture, timestamp_ms: int,
                     vehicle_box: list | None, plate_box: list | None) -> BytesIO | None:
    capture.set(cv2.CAP_PROP_POS_MSEC, max(0, timestamp_ms))
    ok, frame = capture.read()
    if not ok:
        return None
    height, width = frame.shape[:2]
    scale = min(1.0, 640 / width)
    if scale < 1:
        frame = cv2.resize(frame, (round(width * scale), round(height * scale)),
                           interpolation=cv2.INTER_AREA)
    for box, label, color in ((vehicle_box, "VEHICLE", (70, 200, 70)),
                              (plate_box, "PLATE", (0, 200, 255))):
        if not box or len(box) != 4:
            continue
        x1, y1, x2, y2 = (round(float(value) * scale) for value in box)
        x1 = max(0, min(frame.shape[1] - 1, x1))
        x2 = max(0, min(frame.shape[1] - 1, x2))
        y1 = max(0, min(frame.shape[0] - 1, y1))
        y2 = max(0, min(frame.shape[0] - 1, y2))
        if x2 <= x1 or y2 <= y1:
            continue
        cv2.rectangle(frame, (x1, y1), (x2, y2), color, 2)
        cv2.putText(frame, label, (x1, max(14, y1 - 5)), cv2.FONT_HERSHEY_SIMPLEX,
                    .55, color, 2, cv2.LINE_AA)
    ok, encoded = cv2.imencode(".png", frame)
    return BytesIO(encoded.tobytes()) if ok else None


def make_review_workbook(run: dict, rows: list[dict], source_events: list[dict]) -> bytes:
    root = Path(run["run_dir"])
    video_path = Path(run["video_path"])
    plate_path = root / "plate_observations.jsonl"
    plate_by_event = {item["event_id"]: item for item in
                      (load_jsonl(plate_path) if plate_path.is_file() else [])}
    detection_path = root / "detections.jsonl"
    detections = load_jsonl(detection_path) if detection_path.is_file() else []
    event_by_id = {item["event_id"]: item for item in source_events}
    book = Workbook()
    sheet = book.active
    sheet.title = "Ground truth"
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:L{max(1, len(rows) + 1)}"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17613F")
        cell.alignment = Alignment(vertical="center")
    widths = {"A": 18, "B": 15, "C": 15, "D": 16, "E": 12, "F": 20,
              "G": 15, "H": 18, "I": 30, "J": 20, "K": 26, "L": 68}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    capture = cv2.VideoCapture(str(video_path))
    if not capture.isOpened():
        raise ValueError("Không mở được video raw để lấy frame đại diện")
    try:
        for index, row in enumerate(rows, start=2):
            event_id = row["event_id"]
            event = event_by_id.get(event_id)
            sheet.append((event_id, row["start_ms"], row["end_ms"], row["vehicle_type"],
                          row["crossed"], row.get("plate_text") or "none",
                          "none" if row.get("plate_readable") is None else row["plate_readable"],
                          row["condition_status"], ", ".join(row.get("conditions") or []) or "none",
                          event.get("clip_id", "none") if event else "none",
                          run["video_filename"], ""))
            timestamp, vehicle_box, plate_box = _best_frame(
                event, row["start_ms"], row["end_ms"],
                plate_by_event.get(event_id), detections)
            picture = _annotated_image(capture, timestamp, vehicle_box, plate_box)
            if picture:
                excel_image = ExcelImage(picture)
                sheet.add_image(excel_image, f"L{index}")
                sheet.row_dimensions[index].height = excel_image.height * .75 + 8
            else:
                sheet.cell(index, 12, "Không lấy được frame")
            for cell in sheet[index][:11]:
                cell.alignment = Alignment(vertical="center")
    finally:
        capture.release()
    output = BytesIO()
    book.save(output)
    return output.getvalue()
